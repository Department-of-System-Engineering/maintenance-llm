# scripts/generate_synthetic_data.py
"""
Szintetikus demóadatok generálása a projekthez (a valós adatok helyett).

Mit készít:
    1. SQL Server adatbázisok:  ICMMS_INT_DEV (dbo.t_cmms_wo munkalapokkal) és LLM_DEV (üres,
       a vektortáblát a document_processor hozza létre).
    2. Dokumentumok a pdfs/ mappába: karbantartási utasítások PDF-ben + kenési terv Excelben.

Minden adat KITALÁLT. A gépek, hibakódok és dokumentumok egymásra hivatkoznak,
így a dokumentum-, az adatbázis- és a vegyes kérdések is kipróbálhatók.

Futtatás a projekt gyökeréből:
    python -m scripts.generate_synthetic_data                    # minden
    python -m scripts.generate_synthetic_data --only docs        # csak a fájlok
    python -m scripts.generate_synthetic_data --only sql --server localhost
"""

import argparse
import random
from datetime import datetime, timedelta
from pathlib import Path
from urllib.parse import quote_plus

import fitz  # PyMuPDF
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from sqlalchemy import create_engine, text
from sqlalchemy.types import Boolean, DateTime, Integer, NVARCHAR

from scripts.synthetic_docs import DISCLAIMER, DOCUMENTS

SEED = 42
N_WORK_ORDERS = 800
START_DATE = datetime(2025, 1, 2)
END_DATE = datetime(2026, 10, 20)

PROJECT_ROOT = Path(__file__).resolve().parent.parent
DOCS_DIR = PROJECT_ROOT / "pdfs"

ODBC_DRIVER = "ODBC Driver 18 for SQL Server"


# ==========================================================================
# Kitalált üzem: gépek
# ==========================================================================
# (gépszám, géptípus, gépcsoport, részleg, csarnok)
MACHINES = [
    ("#100", "JR20/20 Eszterga", "CNC esztergák", "Forgácsolás (Machining)", "A csarnok"),
    ("#105", "JR20/20 Eszterga", "CNC esztergák", "Forgácsolás (Machining)", "A csarnok"),
    ("#200", "JR20/20 Eszterga", "CNC esztergák", "Forgácsolás (Machining)", "A csarnok"),
    ("#110", "VMC-850 Megmunkáló központ", "Megmunkáló központok", "Forgácsolás (Machining)", "A csarnok"),
    ("#115", "VMC-850 Megmunkáló központ", "Megmunkáló központok", "Forgácsolás (Machining)", "A csarnok"),
    ("#268", "SG-40 Sorjázó gép", "Sorjázó gépek", "Sorjázás és mosás", "A csarnok"),
    ("#270", "SG-40 Sorjázó gép", "Sorjázó gépek", "Sorjázás és mosás", "A csarnok"),
    ("#280", "MW-12 Ipari mosógép", "Mosógépek", "Sorjázás és mosás", "A csarnok"),
    ("#300", "R-6 Robotcella", "Robotcellák", "Szerelés", "B csarnok"),
    ("#310", "PP-25 Pneumatikus prés", "Prések", "Szerelés", "B csarnok"),
    ("#400", "HP-160 Hidraulikus prés", "Prések", "Szerelés", "B csarnok"),
    ("#900", "KS-55 Csavarkompresszor", "Sűrített levegő", "Közműellátás", "Kompresszorház"),
    ("#901", "KS-55 Csavarkompresszor (tartalék)", "Sűrített levegő", "Közműellátás", "Kompresszorház"),
    ("#905", "CH-200 Hűtővíz-előállító", "Hűtővíz", "Közműellátás", "Kompresszorház"),
]

TYPE_OF = {m[0]: m[1] for m in MACHINES}


def machine_type_key(eq_code: str) -> str:
    t = TYPE_OF[eq_code]
    for key in ("JR20/20", "VMC-850", "SG-40", "MW-12", "R-6", "PP-25", "HP-160", "KS-55", "CH-200"):
        if key in t:
            return key
    raise KeyError(t)


# ==========================================================================
# Személyek (kitalált nevek)
# ==========================================================================
TECHNICIANS = [
    (101, "Kovács Péter"), (102, "Nagy Gábor"), (103, "Tóth Zoltán"), (104, "Szabó Attila"),
    (105, "Horváth Bence"), (106, "Varga Csaba"), (107, "Kiss Norbert"), (108, "Molnár Dávid"),
]
NOTIFIERS = [
    (201, "Farkas Anna"), (202, "Balogh Tamás"), (203, "Papp Erika"), (204, "Takács Ádám"),
    (205, "Juhász Réka"), (206, "Lakatos Márk"), (207, "Mészáros Lilla"), (208, "Oláh Krisztián"),
    (209, "Simon Eszter"), (210, "Rácz Viktor"),
]

# ==========================================================================
# Hibaesetek. Mezők:
#   types:   mely géptípusokon fordulhat elő
#   code:    fail_codes
#   lvl:     f_text_lvl1..5 (kategória, részegység, tünet, ok, intézkedés)
#   descr:   wo_descr változatok (bejelentés szövege)
#   sol:     fail_solution változatok
#   dur:     javítási idő percben (min, max)
#   down:    leállással jár-e
#   w:       relatív gyakoriság
# ==========================================================================
FAILURES = [
    dict(types=["JR20/20"], code="KEN-01-01", w=10, dur=(30, 150), down=True,
         lvl=("Kenés", "Főorsó kenés", "Alacsony kenőnyomás (2041)", "Alacsony olajszint", "Olaj utántöltése"),
         descr=["Letiltott a gép főorsó kenés nyomás alacsony hibával (2041).",
                "2041-es riasztás, főorsó kenőnyomás alacsony, gép áll.",
                "Főorsó kenés nyomás alacsony hibaüzenet, operátor szerint kevés az olaj."],
         sol=["Olajszint minimum alatt, {n} l ATF utántöltve. Szivárgás a kenőcső csatlakozásán, csatlakozás meghúzva. Nyomás {p} bar, gép átadva.",
              "ATF olaj utántöltve ({n} l). Nyomás ellenőrizve manométerrel: {p} bar. Próbajáratás rendben."]),
    dict(types=["JR20/20"], code="KEN-01-02", w=7, dur=(60, 180), down=True,
         lvl=("Kenés", "Főorsó kenés", "Alacsony kenőnyomás (2041)", "Eltömődött szűrő", "Szűrőcsere / tisztítás"),
         descr=["Főorsó kenés nyomás alacsony (2041), olajszint rendben.",
                "Ismétlődő 2041-es riasztás, a szivattyú hangosan sípol."],
         sol=["Nyomóoldali szűrőbetét eltömődött (jelző piros), NSZ-10-JR betét cserélve. Nyomás {p} bar.",
              "Szívószűrő (KS-150-JR) erősen szennyezett, kitisztítva. Kavitáció megszűnt, nyomás {p} bar."]),
    dict(types=["JR20/20"], code="KEN-01-03", w=2, dur=(180, 420), down=True,
         lvl=("Kenés", "Főorsó kenés", "Alacsony kenőnyomás (2041)", "Kopott kenőszivattyú", "Szivattyúcsere"),
         descr=["2041 főorsó kenés nyomás alacsony, melegen nem tartja a nyomást.",
                "Főorsó kenőnyomás 0 bar, motor forog."],
         sol=["Szivattyú szállítása 0,7 l/perc (min. 1,2), FSZ-JR-02 szivattyú cserélve. Nyomás {p} bar.",
              "Tengelykapcsoló elnyíródott, tengelykapcsoló cserélve, szivattyú szállítása rendben, nyomás {p} bar."]),
    dict(types=["JR20/20"], code="KEN-01-04", w=2, dur=(30, 90), down=True,
         lvl=("Elektromos", "Főorsó kenés", "Alacsony kenőnyomás (2041)", "Hibás nyomáskapcsoló", "Nyomáskapcsoló csere"),
         descr=["2041 riasztás, de a manométer szerint van nyomás."],
         sol=["PS-41 nyomáskapcsoló hibás, kapcsolási pont elállítódott. Kapcsoló cserélve, bemenet (I12.3) rendben."]),
    dict(types=["JR20/20", "VMC-850"], code="HUT-02-01", w=6, dur=(20, 90), down=True,
         lvl=("Hűtés", "Hűtő-kenő folyadék", "Nincs hűtőfolyadék-áram (3015)", "Eltömődött szívókosár", "Tisztítás"),
         descr=["COOLANT FLOW LOW, 3015-ös riasztás.", "Nem jön a hűtőfolyadék a szerszámhoz."],
         sol=["Hűtőfolyadék-szivattyú szívókosara forgáccsal eltömődött, kitisztítva. Szint utántöltve.",
              "Forgácsszűrő tele, ürítve és tisztítva, áramlás rendben."]),
    dict(types=["JR20/20", "VMC-850"], code="HUT-02-02", w=4, dur=(30, 120), down=False,
         lvl=("Hűtés", "Hűtő-kenő folyadék", "Habzás / szag", "Nem megfelelő koncentráció", "Korrekció / csere"),
         descr=["Emulzió büdös és habzik.", "Hűtő-kenő folyadék koncentráció mérés: {c} %."],
         sol=["Koncentráció {c} %, korrigálva 7 %-ra. pH 8,9.", "pH 7,8, bakteriális fertőzés, teljes emulziócsere az MWI-90105 szerint."]),
    dict(types=["VMC-850"], code="ATC-03-01", w=7, dur=(30, 180), down=True,
         lvl=("Mechanikai", "Szerszámcserélő (ATC)", "Kar nem érte el a végállást (2310)", "Szennyezett közelítéskapcsoló", "Tisztítás, beállítás"),
         descr=["Szerszámcsere közben megállt, 2310-es hiba.", "ATC kar elakadt csere közben."],
         sol=["Közelítéskapcsoló olajos forgáccsal szennyezett, tisztítva, távolság 1,5 mm-re állítva. 10 próbacsere rendben.",
              "Kar kézzel alaphelyzetbe forgatva, M190 referencia, kapcsoló beállítva."]),
    dict(types=["VMC-850"], code="ATC-03-02", w=4, dur=(40, 150), down=True,
         lvl=("Pneumatika", "Szerszámcserélő (ATC)", "Szerszám oldás nem történt meg (2315)", "Alacsony levegőnyomás", "Nyomás helyreállítása"),
         descr=["2315 szerszám oldás hiba.", "Nem engedi el a szerszámot az orsó."],
         sol=["Bemenő levegőnyomás 4,8 bar, FRL szabályzó utánállítva 6 bar-ra. Erősítő olaj utántöltve.",
              "Oldó mágnesszelep tekercse hibás, cserélve."]),
    dict(types=["VMC-850"], code="ATC-03-03", w=2, dur=(120, 300), down=True,
         lvl=("Mechanikai", "Főorsó befogás", "Szerszám nincs megfogva (2331)", "Gyenge tányérrugó", "Rugóköteg csere"),
         descr=["2331 szerszám nincs megfogva, rezeg a megmunkálás."],
         sol=["Befogóerő 6,1 kN (min. 8 kN). Tányérrugó-köteg cserélve, befogóerő 10,2 kN."]),
    dict(types=["SG-40"], code="PNE-04-01", w=8, dur=(30, 90), down=True,
         lvl=("Pneumatika", "Levegőelőkészítő (FRL)", "Pneumatika nyomás alacsony (P-07)", "Hibás nyomásszabályzó", "FRL egység csere"),
         descr=["P-07 pneumatika nyomás alacsony, gép áll.", "Sorjázó levegőelőkészítő szivárog, nyomás nem tartható.",
                "Levegőelőkészítő szűrőháza megrepedt."],
         sol=["FRL egység cserélve (FRL-12-5-SG) az MWI-90102 szerint, nyomás 6,0 bar-ra állítva, szivárgásmentes.",
              "Nyomásszabályzó membrán szakadt, levegőelőkészítő egység komplett csere. Nyomás 6 bar."]),
    dict(types=["SG-40"], code="PNE-04-02", w=4, dur=(15, 45), down=False,
         lvl=("Pneumatika", "Levegőelőkészítő (FRL)", "Víz a pneumatikában", "Kondenzátum nincs leengedve", "Leeresztés, szűrőcsere"),
         descr=["Víz jön a hengereknél.", "Kondenzátum-leeresztő szivárog a sorjázón."],
         sol=["Szűrőpohár tele kondenzátummal, leengedve, szűrőbetét (SZB-5-12) cserélve.",
              "Automata kondenzátum-leeresztő tisztítva."]),
    dict(types=["SG-40"], code="MEC-05-01", w=6, dur=(45, 120), down=True,
         lvl=("Mechanikai", "Sorjázó kefe", "Kopott kefe / gyenge sorjázás", "Kefe kopás", "Kefecsere"),
         descr=["Sorja marad a darabon, kefe kopott.", "Minőség jelezte: sorjás alkatrészek."],
         sol=["Kefék cserélve (2 db), fogás beállítva.", "Kefe kopott, cserélve, ráhagyás 0,5 mm-re állítva."]),
    dict(types=["MW-12"], code="HUT-06-01", w=4, dur=(30, 120), down=True,
         lvl=("Hűtés", "Mosófolyadék fűtés", "Hőmérséklet nem éri el az alapjelet", "Fűtőpatron hiba", "Fűtőpatron csere"),
         descr=["Mosógép nem melegszik fel 60 fokra.", "Mosófolyadék hideg, darabok olajosak maradnak."],
         sol=["Egyik fűtőpatron szakadt (mért ellenállás végtelen), cserélve.", "Hőfokszabályzó hibás, cserélve."]),
    dict(types=["MW-12"], code="MEC-06-02", w=3, dur=(30, 90), down=True,
         lvl=("Mechanikai", "Mosószivattyú", "Gyenge mosónyomás", "Eltömődött fúvókák", "Fúvókatisztítás"),
         descr=["Gyenge a mosás, fúvókák nem spriccelnek."],
         sol=["Fúvókák vízkővel eltömődve, kiszerelve, tisztítva. Szűrőkosár ürítve."]),
    dict(types=["R-6"], code="ROB-07-01", w=6, dur=(30, 150), down=True,
         lvl=("Mechanikai", "Robot", "Ütközés észlelve (SRVO-050)", "Rosszul pozícionált munkadarab", "Kiszabadítás, ellenőrzés"),
         descr=["SRVO-050 ütközés, robot áll.", "Robot nekiment a készüléknek."],
         sol=["Robot kiszabadítva T1 módban az MWI-90107 szerint. REF_CHECK eltérés 0,2 mm, rendben. Megfogó ép.",
              "Ütközés után megfogó ujj sérült, cserélve. Referencia ellenőrzés rendben (0,3 mm)."]),
    dict(types=["R-6"], code="ROB-07-02", w=3, dur=(20, 60), down=True,
         lvl=("Pneumatika", "Megfogó", "Megfogó nem zár", "Sérült pneumatikus cső", "Csőcsere"),
         descr=["Megfogó nem fogja meg a darabot, leejti."],
         sol=["Megfogó pneumatikus csöve kidörzsölődött, cserélve. Nyomás 5 bar."]),
    dict(types=["PP-25"], code="PNE-08-01", w=3, dur=(30, 90), down=True,
         lvl=("Pneumatika", "Préshenger", "Lassú / gyenge préselés", "Tömítetlen henger", "Tömítéscsere"),
         descr=["Prés lassan jár le, gyenge.", "Préshenger szivárog."],
         sol=["Henger dugattyútömítése kopott, tömítéskészlet cserélve.", "Levegőelőkészítő szabályzó elállítódott, 6 bar-ra állítva."]),
    dict(types=["HP-160"], code="HID-09-01", w=6, dur=(30, 180), down=True,
         lvl=("Hidraulika", "Olajhűtés", "Olajhőmérséklet magas (H-12/H-13)", "Nincs hűtővíz", "Hűtővíz helyreállítása"),
         descr=["H-13 olajhőmérséklet magas, prés leállt.", "H-12 figyelmeztetés, olaj 63 fok."],
         sol=["Hűtővíz-szelep zárva maradt a CH-200 karbantartás után, kinyitva. Olaj lehűlt 48 °C-ra.",
              "Hőcserélő vízoldala vízköves, vízkőmentesítve. Átfolyás 28 l/perc."]),
    dict(types=["HP-160"], code="HID-09-02", w=4, dur=(60, 240), down=True,
         lvl=("Hidraulika", "Hidraulika rendszer", "Olajszivárgás", "Sérült tömlő / tömítés", "Tömlő- vagy tömítéscsere"),
         descr=["Olajfolt a prés alatt.", "Hidraulika tömlő szivárog a hengernél."],
         sol=["Hengerhez menő nagynyomású tömlő cserélve, {n} l HLP 46 utántöltve.", "Szeleptömb O-gyűrű cserélve, szivárgás megszűnt."]),
    dict(types=["KS-55"], code="KOM-10-01", w=5, dur=(30, 120), down=True,
         lvl=("Hűtés", "Kompresszor hűtő", "Csavarblokk hőmérséklet magas (E04)", "Elszennyeződött hűtő", "Hűtő tisztítása"),
         descr=["Kompresszor E04 hibával leállt.", "Kompresszor túlmelegedett, hálózati nyomás leesett."],
         sol=["Hűtő porral eltömődve, kifújva. Tartalék kompresszor indítva a javítás idejére. Hőmérséklet 84 °C.",
              "Olajszint alacsony, {n} l KO-46 utántöltve, hűtő tisztítva."]),
    dict(types=["KS-55"], code="KOM-10-02", w=2, dur=(20, 60), down=False,
         lvl=("Elektromos", "Kondenzátum-leeresztő", "Kondenzátum-leeresztő hiba (E12)", "Eltömődött szelep", "Tisztítás"),
         descr=["E12 kondenzátum-leeresztő hiba."],
         sol=["Leeresztő szelep szennyezett, kitisztítva, mágnestekercs rendben."]),
    dict(types=["CH-200"], code="HUT-11-01", w=3, dur=(60, 240), down=True,
         lvl=("Hűtés", "Hűtőkör", "Magas hűtővíz-hőmérséklet", "Alacsony hűtőközeg-töltet", "Szivárgáskeresés, töltés"),
         descr=["Hűtővíz előremenő 22 fok, hűtő nem hűt rendesen."],
         sol=["Hűtőközeg-szivárgás a kondenzátornál, külső szerviz javította és utántöltötte."]),
    dict(types=["JR20/20", "VMC-850", "SG-40", "MW-12", "R-6", "PP-25", "HP-160", "KS-55", "CH-200"],
         code="ELE-12-01", w=6, dur=(20, 120), down=True,
         lvl=("Elektromos", "Kapcsolószekrény", "Hőkioldás / túlmelegedés", "Eltömődött szűrő a szekrényventilátoron", "Szűrőcsere, ventilátor csere"),
         descr=["Gép leállt, kapcsolószekrényben hőkioldás.", "Kapcsolószekrény ventilátor nem forog."],
         sol=["Szekrényventilátor szűrőbetétje eltömődött, cserélve.", "Ventilátor csapágya szorult, ventilátor cserélve."]),
    dict(types=["JR20/20", "VMC-850", "SG-40", "MW-12", "R-6", "PP-25", "HP-160"],
         code="ELE-12-02", w=4, dur=(10, 60), down=True,
         lvl=("Elektromos", "Biztonsági kör", "Vészleállító kör megszakadt", "Hibás ajtókapcsoló", "Kapcsoló csere / beállítás"),
         descr=["Nem indul, ajtó nyitva hibát jelez.", "Biztonsági ajtó reteszelés hiba."],
         sol=["Ajtó biztonsági kapcsoló elállítódott, beállítva.", "Biztonsági ajtókapcsoló cserélve."]),
]

# Megelőző karbantartások: (géptípusok, leírás, időtartam percben, f_text_lvl1, részegység)
PREVENTIVE = [
    (["JR20/20"], "Havi kenés és ellenőrzés a KP-2026 kenési terv szerint (főorsó csapágyzsírzás, kenőolajszint).", (45, 90), "Kenés", "Főorsó kenés"),
    (["JR20/20"], "Negyedéves szűrőcsere főorsó kenésen (NSZ-10-JR) – MWI-90101.", (30, 60), "Kenés", "Főorsó kenés"),
    (["JR20/20", "VMC-850"], "Hűtő-kenő folyadék koncentráció és pH ellenőrzés – MWI-90105.", (15, 30), "Hűtés", "Hűtő-kenő folyadék"),
    (["VMC-850"], "Havi ATC karbantartás: közelítéskapcsolók tisztítása, erősítő olajszint – MWI-90106.", (40, 80), "Mechanikai", "Szerszámcserélő (ATC)"),
    (["SG-40"], "Negyedéves FRL szűrőbetét csere (SZB-5-12) – MWI-90102.", (20, 40), "Pneumatika", "Levegőelőkészítő (FRL)"),
    (["MW-12"], "Havi fúvóka- és szűrőkosár tisztítás.", (30, 60), "Mechanikai", "Mosószivattyú"),
    (["R-6"], "Negyedéves robot karbantartás: REF_CHECK, kábelköteg, megfogó ellenőrzés.", (60, 120), "Mechanikai", "Robot"),
    (["PP-25"], "Havi pneumatika ellenőrzés, szivárgáskeresés.", (20, 40), "Pneumatika", "Préshenger"),
    (["HP-160"], "Féléves hidraulika szűrőcsere (HSZ-10-160, LSZ-HP-01) – MWI-90103.", (90, 150), "Hidraulika", "Hidraulika rendszer"),
    (["KS-55"], "2000 üzemórás karbantartás: levegőszűrő csere (LSZ-55), hűtő tisztítás – MWI-90104.", (60, 120), "Hűtés", "Kompresszor hűtő"),
    (["CH-200"], "Negyedéves hűtő ellenőrzés: kondenzátor tisztítás, nyomások rögzítése.", (60, 120), "Hűtés", "Hűtőkör"),
]

CATEGORIES = {
    "BD": "Hibaelhárítás",
    "PM": "Megelőző karbantartás",
    "CM": "Javító karbantartás",
    "INS": "Ellenőrzés",
}


def _random_dt(rng: random.Random) -> datetime:
    span = (END_DATE - START_DATE).total_seconds()
    dt = START_DATE + timedelta(seconds=rng.uniform(0, span))
    return dt.replace(second=0, microsecond=0)


def _fill(template: str, rng: random.Random) -> str:
    return template.format(
        n=rng.choice([2, 3, 4, 5]),
        p=f"{rng.uniform(2.6, 3.4):.1f}".replace(".", ","),
        c=f"{rng.choice([4.2, 4.8, 5.1, 9.2, 9.8, 10.5])}".replace(".", ","),
    )


def _path_columns(eq_code: str) -> dict:
    m = next(x for x in MACHINES if x[0] == eq_code)
    lvls = ["Mintaüzem", m[4], m[3], m[2], f"{m[0]} {m[1]}"]
    row = {f"path_lvl{i + 1}": v for i, v in enumerate(lvls)}
    row["eq_full_path"] = " / ".join(lvls)
    return row


def generate_work_orders(n: int = N_WORK_ORDERS, seed: int = SEED) -> pd.DataFrame:
    rng = random.Random(seed)
    rows = []
    eq_codes = [m[0] for m in MACHINES]
    # A forgácsoló gépek és a sorjázók többet hibáznak — ez adja a "kiugró" gépeket
    eq_weights = [1.6 if machine_type_key(e) in ("JR20/20", "VMC-850", "SG-40") else 1.0 for e in eq_codes]
    eq_weights[eq_codes.index("#268")] = 2.4
    eq_weights[eq_codes.index("#901")] = 0.15

    for i in range(n):
        eq = rng.choices(eq_codes, weights=eq_weights)[0]
        tkey = machine_type_key(eq)
        notifier_id, notifier = rng.choice(NOTIFIERS)
        resp_id, resp = rng.choice(TECHNICIANS)
        rec = _random_dt(rng)

        kind = rng.choices(["BD", "PM", "CM", "INS"], weights=[62, 25, 8, 5])[0]
        pm_options = [p for p in PREVENTIVE if tkey in p[0]]
        if kind in ("PM", "INS") and pm_options:
            pm = rng.choice(pm_options)
            dur = rng.randint(*pm[2])
            wo_descr, fail_solution = pm[1], "Elvégezve, eltérés nem volt." if rng.random() < 0.85 else \
                "Elvégezve. Eltérés: kopás jelei, javítás tervezve."
            lvl = (pm[3], pm[4], None, None, None)
            fail_code, fail_text = None, None
            priority, is_down = 4, rng.random() < 0.5
            due = rec.replace(hour=6, minute=0) + timedelta(days=rng.randint(1, 7))
            start = due + timedelta(minutes=rng.randint(0, 600))
        else:
            if kind in ("PM", "INS"):
                kind = "BD"
            fails = [f for f in FAILURES if tkey in f["types"]]
            f = rng.choices(fails, weights=[x["w"] for x in fails])[0]
            dur = rng.randint(*f["dur"])
            wo_descr = _fill(rng.choice(f["descr"]), rng)
            fail_solution = _fill(rng.choice(f["sol"]), rng)
            lvl = f["lvl"]
            fail_code, fail_text = f["code"], f"{f['lvl'][2]} – {f['lvl'][3]}"
            is_down = f["down"] if kind == "BD" else rng.random() < 0.3
            priority = (rng.choice([1, 1, 2]) if is_down else rng.choice([2, 3])) if kind == "BD" else 3
            due = rec + timedelta(hours={1: 4, 2: 24, 3: 72, 4: 168}[priority])
            start = rec + timedelta(minutes=rng.randint(5, 50) if kind == "BD" else rng.randint(120, 4000))

        end = start + timedelta(minutes=dur)
        done = end < END_DATE - timedelta(days=3) or rng.random() < 0.4
        techs = 2 if dur > 150 and rng.random() < 0.6 else 1

        eq_stop = eq_start = None
        if is_down:
            eq_stop = rec - timedelta(minutes=rng.randint(1, 20)) if kind == "BD" else start
            eq_start = end + timedelta(minutes=rng.randint(0, 15)) if done else None

        row = {
            "seq_no": 250000 + i,
            "rec_date": rec,
            "wo_start": start,
            "wo_end": end if done else None,
            "wo_due_date": due,
            "eq_stop": eq_stop,
            "eq_start": eq_start,
            "fail_codes": fail_code,
            "fail_codes_text": fail_text,
            "fail_solution": fail_solution if done else None,
            "wo_descr": wo_descr,
            "priority": priority,
            "resp_name": resp,
            "notifier_name": notifier,
            "wo_cat_name": CATEGORIES[kind],
            "wo_cat_code": kind,
            "is_bd": kind == "BD",
            "eq_is_down": bool(is_down),
            "eq_short_code": eq,
            "wo_is_done": bool(done),
            "sum_work_hour_in_sec": dur * 60 * techs if done else None,
            "resp_id": resp_id,
            "notifier_id": notifier_id,
            **{f"f_text_lvl{k + 1}": lvl[k] for k in range(5)},
            **_path_columns(eq),
            "wo_duration_minutes": dur if done else None,
            "eq_duration_minutes": int((eq_start - eq_stop).total_seconds() // 60) if eq_stop and eq_start else None,
        }
        rows.append(row)

    df = pd.DataFrame(rows).sort_values("rec_date").reset_index(drop=True)
    df["seq_no"] = range(250001, 250001 + len(df))
    return df


# ==========================================================================
# SQL Server
# ==========================================================================
T_CMMS_WO_DDL = """
CREATE TABLE dbo.t_cmms_wo (
    seq_no               INT            NOT NULL PRIMARY KEY,
    rec_date             DATETIME       NOT NULL,
    wo_start             DATETIME       NULL,
    wo_end               DATETIME       NULL,
    wo_due_date          DATETIME       NULL,
    eq_stop              DATETIME       NULL,
    eq_start             DATETIME       NULL,
    fail_codes           NVARCHAR(50)   NULL,
    fail_codes_text      NVARCHAR(255)  NULL,
    fail_solution        NVARCHAR(1000) NULL,
    wo_descr             NVARCHAR(1000) NULL,
    priority             INT            NULL,
    resp_name            NVARCHAR(100)  NULL,
    notifier_name        NVARCHAR(100)  NULL,
    wo_cat_name          NVARCHAR(100)  NULL,
    wo_cat_code          NVARCHAR(20)   NULL,
    is_bd                BIT            NULL,
    eq_is_down           BIT            NULL,
    eq_short_code        NVARCHAR(50)   NULL,
    eq_full_path         NVARCHAR(500)  NULL,
    wo_is_done           BIT            NULL,
    sum_work_hour_in_sec INT            NULL,
    resp_id              INT            NULL,
    notifier_id          INT            NULL,
    f_text_lvl1          NVARCHAR(255)  NULL,
    f_text_lvl2          NVARCHAR(255)  NULL,
    f_text_lvl3          NVARCHAR(255)  NULL,
    f_text_lvl4          NVARCHAR(255)  NULL,
    f_text_lvl5          NVARCHAR(255)  NULL,
    path_lvl1            NVARCHAR(255)  NULL,
    path_lvl2            NVARCHAR(255)  NULL,
    path_lvl3            NVARCHAR(255)  NULL,
    path_lvl4            NVARCHAR(255)  NULL,
    path_lvl5            NVARCHAR(255)  NULL,
    wo_duration_minutes  INT            NULL,
    eq_duration_minutes  INT            NULL
)
"""


def build_uri(server: str, database: str) -> str:
    return (
        f"mssql+pyodbc://@{server}/{database}"
        f"?driver={quote_plus(ODBC_DRIVER)}&trusted_connection=yes&TrustServerCertificate=yes"
    )


def create_databases(server: str, names: list[str]) -> None:
    master = create_engine(build_uri(server, "master"), isolation_level="AUTOCOMMIT")
    with master.connect() as conn:
        for name in names:
            conn.execute(text(f"IF DB_ID(N'{name}') IS NULL CREATE DATABASE [{name}]"))
            print(f"Adatbázis kész: {name}")


def load_work_orders(server: str, df: pd.DataFrame) -> None:
    engine = create_engine(build_uri(server, "ICMMS_INT_DEV"), fast_executemany=True)
    with engine.begin() as conn:
        conn.execute(text("IF OBJECT_ID('dbo.t_cmms_wo') IS NOT NULL DROP TABLE dbo.t_cmms_wo"))
        conn.execute(text(T_CMMS_WO_DDL))
        df.to_sql("t_cmms_wo", conn, schema="dbo", if_exists="append", index=False, chunksize=500)
    print(f"dbo.t_cmms_wo feltöltve: {len(df)} munkalap")


# ==========================================================================
# Dokumentumok
# ==========================================================================
# A MuPDF beépített sans-serif betűtípusa: ezzel a pypdf tisztán (valódi kötőjellel,
# szóközzel) nyeri ki a szöveget. TTF betűtípussal a '-' lágy kötőjelként jönne vissza.
PDF_CSS = """
body { font-family: sans-serif; font-size: 10pt; line-height: 1.35; }
h1 { font-size: 15pt; margin-bottom: 4pt; }
h2 { font-size: 12pt; margin-top: 10pt; margin-bottom: 3pt; }
p.disclaimer { font-size: 8pt; color: #a00; }
table { border-collapse: collapse; width: 100%; }
th, td { border: 1px solid #555; padding: 2pt 4pt; font-size: 9pt; vertical-align: top; }
th { background-color: #ddd; }
"""


def write_pdf(path: Path, title: str, body_html: str) -> None:
    html = f"<p class='disclaimer'>{DISCLAIMER}</p><h1>{title}</h1>{body_html}"
    story = fitz.Story(html=html, user_css=PDF_CSS)
    writer = fitz.DocumentWriter(str(path))
    mediabox = fitz.paper_rect("a4")
    where = mediabox + (50, 50, -50, -50)
    more = True
    while more:
        device = writer.begin_page(mediabox)
        more, _ = story.place(where)
        story.draw(device)
        writer.end_page()
    writer.close()


LUBRICATION_PLAN = [
    # gép, gépszámok, kenési/mérési pont, kenőanyag, mennyiség, intervallum, felelős, idő (perc)
    ("JR20/20 Eszterga", "#100; #105; #200", "Főorsó kenőolajszint", "ATF (Dexron III)", "20 l tartály, min. jel felett", "Műszakonként", "Operátor", 1),
    ("JR20/20 Eszterga", "#100; #105; #200", "Főorsó motor csapágyzsírzás", "LT/WR EP 2 zsír", "2 nyomás / zsírzógomb", "Havi", "Karbantartó", 5),
    ("JR20/20 Eszterga", "#100; #105; #200", "Főorsó kenés nyomóoldali szűrő (NSZ-10-JR)", "-", "1 db betét", "Negyedéves", "Karbantartó", 20),
    ("JR20/20 Eszterga", "#100; #105; #200", "Szánvezeték központi kenés", "Vezetékpályaolaj ISO VG 68", "2 l tartály", "Heti", "Operátor", 2),
    ("JR20/20 Eszterga", "#100; #105; #200", "Hűtő-kenő folyadék koncentráció", "Emulzió 6–8 %", "180 l", "Műszakonként", "Operátor", 3),
    ("VMC-850 Megmunkáló központ", "#110; #115", "ATC erősítő olajszint", "HLP 32", "0,5 l", "Havi", "Karbantartó", 5),
    ("VMC-850 Megmunkáló központ", "#110; #115", "Szánvezeték központi kenés", "Vezetékpályaolaj ISO VG 68", "3 l tartály", "Heti", "Operátor", 2),
    ("VMC-850 Megmunkáló központ", "#110; #115", "Hűtő-kenő folyadék koncentráció", "Emulzió 6–8 %", "300 l", "Műszakonként", "Operátor", 3),
    ("SG-40 Sorjázó gép", "#268; #270", "FRL olajozó szintje", "Pneumatika olaj ISO VG 32", "0,1 l", "Heti", "Operátor", 1),
    ("SG-40 Sorjázó gép", "#268; #270", "Kefe-tengely csapágyak", "LT/WR EP 2 zsír", "1 nyomás / csapágy", "Havi", "Karbantartó", 10),
    ("SG-40 Sorjázó gép", "#268; #270", "FRL szűrőbetét (SZB-5-12)", "-", "1 db", "Negyedéves", "Karbantartó", 15),
    ("R-6 Robotcella", "#300", "Robot hajtóművek (J1–J3)", "Robot hajtómű zsír RZ-1", "Gyártói előírás szerint", "Éves", "Karbantartó", 120),
    ("HP-160 Hidraulikus prés", "#400", "Hidraulikaolaj szint és hőmérséklet", "HLP 46", "160 l tartály", "Napi", "Operátor", 2),
    ("HP-160 Hidraulikus prés", "#400", "Visszatérő ági szűrő (HSZ-10-160)", "-", "1 db", "Féléves", "Karbantartó", 45),
    ("HP-160 Hidraulikus prés", "#400", "Préskos vezetékek", "LT/WR EP 2 zsír", "4 zsírzógomb", "Heti", "Karbantartó", 10),
    ("KS-55 Csavarkompresszor", "#900; #901", "Kompresszorolaj szint", "KO-46 szintetikus", "18 l töltet", "Napi", "Közmű karbantartó", 2),
    ("KS-55 Csavarkompresszor", "#900; #901", "Olajszűrő és olajcsere (OSZ-55)", "KO-46 szintetikus", "18 l + 1 db szűrő", "4000 üzemóra", "Karbantartó", 90),
]


def write_lubrication_plan(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Kenési terv"
    ws.append([DISCLAIMER])
    ws.append(["KP-2026 – Kenési és ellenőrzési terv, Mintaüzem"])
    ws.append([])
    header = ["Gép", "Gépszámok", "Kenési pont / Mérési pont", "Kenőanyag típus",
              "Kenőanyag mennyiség", "Intervallum", "Felelős", "Időtartam (perc)"]
    ws.append(header)
    for row in LUBRICATION_PLAN:
        ws.append(list(row))

    ws["A1"].font = Font(color="AA0000", italic=True, size=9)
    ws["A2"].font = Font(bold=True, size=13)
    for cell in ws[4]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="DDDDDD")
    for col, width in zip("ABCDEFGH", (28, 18, 42, 28, 26, 16, 18, 10)):
        ws.column_dimensions[col].width = width
    for r in ws.iter_rows(min_row=5):
        for c in r:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    wb.save(path)


def write_documents() -> None:
    DOCS_DIR.mkdir(exist_ok=True)
    for file_name, (title, body) in DOCUMENTS.items():
        write_pdf(DOCS_DIR / file_name, title, body)
        print(f"PDF: {file_name}")
    xlsx = DOCS_DIR / "KP-2026_Kenesi_es_ellenorzesi_terv.xlsx"
    write_lubrication_plan(xlsx)
    print(f"Excel: {xlsx.name}")


def main() -> None:
    parser = argparse.ArgumentParser(description="Szintetikus demóadatok generálása")
    parser.add_argument("--only", choices=["docs", "sql"], help="csak az egyik részt generálja")
    parser.add_argument("--server", default="localhost", help="SQL Server példány (Windows auth)")
    args = parser.parse_args()

    if args.only in (None, "docs"):
        write_documents()
    if args.only in (None, "sql"):
        create_databases(args.server, ["ICMMS_INT_DEV", "LLM_DEV"])
        load_work_orders(args.server, generate_work_orders())


if __name__ == "__main__":
    main()
